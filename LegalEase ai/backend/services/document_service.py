import base64
from io import BytesIO

from docx import Document
from fpdf import FPDF


def make_txt(text: str) -> bytes:
    return text.encode("utf-8")


def make_docx(
    text: str,
    document_type: str,
    terms: list[str],
    logo_base64: str | None,
) -> bytes:
    document = Document()
    if logo_base64:
        document.add_picture(BytesIO(base64.b64decode(logo_base64)), width=None)
    document.add_heading(document_type, level=1)
    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)
    return _document_bytes(document)


def make_pdf(
    text: str,
    document_type: str,
    terms: list[str],
    logo_base64: str | None,
) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 10, document_type)
    pdf.ln(4)
    pdf.set_font("Helvetica", size=11)
    for paragraph in text.splitlines() or [""]:
        if paragraph:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(
                0,
                7,
                paragraph.encode("latin-1", "replace").decode("latin-1"),
            )
        else:
            pdf.ln(7)
    return bytes(pdf.output())


def _document_bytes(document: Document) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()